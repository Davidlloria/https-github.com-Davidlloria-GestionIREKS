import os
from types import SimpleNamespace

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtCore import QPoint, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from app.ui.widgets.product_consumers_dialog import ProductConsumersDialog
from app.services.product_consumers_export import export_excel, export_pdf


def make_dialog():
    app = QApplication.instance() or QApplication([])
    rows = [
        SimpleNamespace(cliente_codigo="100", cliente_nombre="Zeta", kg_prev=1000, euros_prev=12000, kg_curr=9, euros_curr=108),
        SimpleNamespace(cliente_codigo="2", cliente_nombre="Alfa", kg_prev=2, euros_prev=24, kg_curr=100, euros_curr=1200),
        SimpleNamespace(cliente_codigo="30", cliente_nombre="Beta", kg_prev=0, euros_prev=0, kg_curr=20, euros_curr=240),
    ]
    return app, ProductConsumersDialog(2026, "38790", "REX SARRACENO", rows)


def test_sort_all_columns_keeps_records_and_totals_together():
    app, dialog = make_dialog()
    original = {row[0]: row for row in dialog.ordered_rows()}
    totals = dialog.totals[:]
    for col in range(8):
        for direction in (Qt.SortOrder.AscendingOrder, Qt.SortOrder.DescendingOrder):
            dialog.table.sortItems(col, direction)
            rows = dialog.ordered_rows()
            key = (lambda row: int(row[col])) if col == 0 else (lambda row: row[col])
            assert rows == sorted(original.values(), key=key, reverse=direction == Qt.SortOrder.DescendingOrder)
            assert all(row == original[row[0]] for row in rows)
            assert dialog.totals == totals
    assert dialog.table.horizontalHeaderItem(6).text() == "Δ Kilos"
    assert dialog.table.horizontalHeaderItem(7).text() == "Δ €"
    dialog.close()


def test_totals_follow_column_resize_and_horizontal_scroll():
    app, dialog = make_dialog()
    dialog.show()
    app.processEvents()
    dialog.resize(1050, 580)
    dialog.table.setColumnWidth(3, 190)
    app.processEvents()
    app.processEvents()
    dialog.table.horizontalScrollBar().setValue(180)
    app.processEvents()
    for col in range(8):
        assert dialog.footer.columnWidth(col) == dialog.table.columnWidth(col)
        assert dialog.footer.columnViewportPosition(col) == dialog.table.columnViewportPosition(col)
    assert dialog.footer.viewport().width() == dialog.table.viewport().width()
    dialog.close()


def test_export_contains_all_sorted_rows_numeric_values_and_totals(tmp_path):
    app, dialog = make_dialog()
    dialog.table.sortItems(0, Qt.SortOrder.AscendingOrder)
    rows = dialog.ordered_rows()
    rows[0][1] = '=Cliente con nombre muy largo & hijos <Norte> ' * 3
    target = tmp_path / "clientes.xlsx"
    export_excel(target, dialog.subtitle, dialog.export_headers, rows, dialog.totals)
    sheet = load_workbook(target).active
    assert [sheet.cell(row, 1).value for row in range(4, 7)] == ["2", "30", "100"]
    assert sheet.cell(4, 2).data_type == "s"
    assert sheet.cell(4, 5).value == 100
    assert [sheet.cell(7, col).value for col in range(3, 9)] == dialog.totals
    assert sheet.auto_filter.ref == "A3:H6"
    pdf = tmp_path / "clientes.pdf"
    export_pdf(pdf, dialog.subtitle, dialog.export_headers, rows * 30, [value * 30 for value in dialog.totals])
    import fitz
    document = fitz.open(pdf)
    assert len(document) > 1
    text = "".join(page.get_text() for page in document)
    assert text.count("Zeta") == 30
    assert "Totales generales" in text
    assert all("2025" in page.get_text() for page in document)
    document.close()
    dialog.close()


def test_empty_table_disables_exports():
    app = QApplication.instance() or QApplication([])
    dialog = ProductConsumersDialog(2026, "X", "Vacío", [])
    assert dialog.ordered_rows() == []
    assert all(not button.isEnabled() for button in dialog.export_buttons)
    dialog.close()


def test_click_header_toggles_numeric_order():
    app, dialog = make_dialog()
    dialog.show()
    app.processEvents()
    header = dialog.table.horizontalHeader()
    point = QPoint(header.sectionViewportPosition(4) + 30, 50)
    QTest.mouseClick(header.viewport(), Qt.MouseButton.LeftButton, pos=point)
    assert [row[4] for row in dialog.ordered_rows()] == [9, 20, 100]
    QTest.mouseClick(header.viewport(), Qt.MouseButton.LeftButton, pos=point)
    assert [row[4] for row in dialog.ordered_rows()] == [100, 20, 9]
    dialog.close()


def test_export_action_handles_cancel_success_and_failure(monkeypatch, tmp_path):
    from app.ui.widgets import product_consumers_dialog as module
    app, dialog = make_dialog()
    calls = []
    notices = []
    monkeypatch.setattr(module, "export_excel", lambda *args: calls.append(args))
    monkeypatch.setattr(module.QMessageBox, "information", lambda *args: notices.append("ok"))
    monkeypatch.setattr(module.QMessageBox, "warning", lambda *args: notices.append("error"))
    monkeypatch.setattr(module.QFileDialog, "getSaveFileName", lambda *args: ("", ""))
    dialog._export("xlsx")
    assert not calls
    target = str(tmp_path / "clientes")
    monkeypatch.setattr(module.QFileDialog, "getSaveFileName", lambda *args: (target, ""))
    dialog.table.sortItems(6, Qt.SortOrder.AscendingOrder)
    dialog._export("xlsx")
    assert calls[0][0] == target + ".xlsx"
    assert calls[0][3] == dialog.ordered_rows()
    assert notices == ["ok"]
    def fail(*args):
        raise PermissionError("Archivo abierto")
    monkeypatch.setattr(module, "export_excel", fail)
    dialog._export("xlsx")
    assert notices == ["ok", "error"]
    dialog.close()
