from copy import deepcopy
from dataclasses import asdict
from decimal import Decimal
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from sqlmodel import SQLModel, Session, create_engine
from sqlalchemy.pool import StaticPool

from app.models import Cliente, Distribuidor, Familia, Fabricante, IngredienteIreks, VentaMensualRaw
from app.services.objectives_service import Objective, ObjectivesService, score, fmt


@pytest.mark.parametrize("target,maximum,previous,current,rule,expected", [
    (3, 6, 20590, 20985, "growth", "3"),
    (2, 6, 74253.3, 78526.8, "growth", "6"),
    (2, 7, 131672.5, 129522.5, "growth", "0"),
    (100, 4, 18, 29, "growth", "2"),
    (3, 5, 267327.1, 268816.2, "quarter", "0.75"),
    (0, 2, 15531, 14934.5, "no_charge", "2"),
    (0, 2, 100, 100, "no_charge", "0"),
    (0, 2, 100, 101, "no_charge", "0"),
    (0, 5, 100, 100, "growth", "0"),
    (0, 5, 100, 101, "growth", "5"),
    (0, 5, 100, 99, "growth", "0"),
    (3, 6, 0, 10, "growth", "6"),
    (3, 6, 0, 0, "growth", "0"),
    (3, 6, 100, -20, "growth", "0"),
])
def test_workbook_scoring_boundaries(target, maximum, previous, current, rule, expected):
    objective = Objective("x", "Prueba", 0, target, maximum, rule=rule)
    assert score(objective, previous, current, 1000, 7) == Decimal(expected)


@pytest.mark.parametrize("month,expected", [(1, ".4"), (3, "1.2"), (7, "2.4"), (8, "2.8"), (12, "4")])
def test_fixed_launch_calendar_rule(month, expected):
    o = Objective("x", "Lanzamiento", 1, 1500, 4, target_type="kg")
    assert score(o, 0, 687.5, 87.5, month) == Decimal(expected)
    assert score(o, 0, 0, 87.5, month) == 0


def test_official_july_total():
    # Independent reference inputs from the July workbook, not a live DB fixture.
    general = [(6,3,20590,20985),(7,2,131672.5,129522.5),(6,2,74253.3,78526.8),
        (6,1,9632.5,8475.5),(5,5,1431,1245.5),(5,3,2634,1324),(5,0,1015,929),
        (5,4,11156.3,10043.4),(5,5,20458.19,19402.53),(4,100,300,125),
        (4,9,13075,12700),(4,100,18,29),(11,2,230704.6,221938.7),
        (2,10,12400,12425),(4,17,24222.5,34452.5)]
    total = sum(score(Objective("x", "x", 0, target, maximum), prev, curr, 1, 7) for maximum,target,prev,curr in general)
    total += score(Objective("x", "x", 1, 1500, 4, target_type="kg"),0,687.5,87.5,7)
    total += score(Objective("x", "x", 2, 0, 2, rule="no_charge"),15531,14934.5,26598.5,7)
    total += score(Objective("x", "x", 3, 3, 5, rule="quarter"),267327.1,268816.2,469285.6,7)
    assert total == Decimal("20.15")
    assert fmt(Decimal("2512.125")) == "2.512,13"


@pytest.fixture
def service(tmp_path):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    SQLModel.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(Fabricante(fabricante_id="m", fabricante_nombre="IREKS"))
        session.add(Familia(articulo_familia_id="f", fabricante_id="m", articulo_familia_nombre="MEJORANTES"))
        for code, (key, name) in enumerate((("igsa", "IGSA"), ("cadelsa", "CADELSA LZA"), ("hermanos", "Panadería Hermanos Rodríguez"), ("baker", "Backer Las Arenas Sin Gluten"), ("other", "Otra zona")), 1):
            session.add(Cliente(cliente_id=key, cliente_codigo=code, cliente_nombre_comercial=name))
        session.add(Distribuidor(distribuidor_id="igsa-old", distribuidor_nombre_comercial="IGSA"))
        for key, name in (("p", "MEJORANTE"), ("launch", "MELLA CARROT")):
            session.add(IngredienteIreks(articulo_id=key, articulo_referencia=key, articulo_descripcion=name, articulo_familia_id="f", fabricante_id="m"))
        session.commit()
    result = ObjectivesService(engine, tmp_path / "objectives.json")
    yield result
    engine.dispose()


def add_sales(service, *rows):
    with Session(service.engine) as session:
        for party, period, product, kilos, sc, euros, source in rows:
            session.add(VentaMensualRaw(cliente_id=party, periodo=period, articulo_id=product, articulo_codigo_origen=product, venta_kilos=kilos, venta_kilos_sc=sc, venta_euros=euros, fuente=source))
        session.commit()


def compact_campaign(service):
    c = service.default_campaign()
    objectives = [Objective("total", "Volumen total", 3, 3, 5, rule="quarter"),
        Objective("familia", "Mejorantes", 0, 3, 6, dimension="family", members=["f"]),
        Objective("launch", "Lanzamiento", 1, 1500, 4, target_type="kg", dimension="products", members=["launch"]),
        Objective("sc", "Sin cargo", 2, 0, 2, unit="sc", rule="no_charge"),
        Objective("eur", "Euros", 0, 5, 5, unit="eur"),
        Objective("pending", "Panesa", 2, 4, 3, status="pending"),
        Objective("ten", "Ten", 1, 750, 4, target_type="kg", status="unmarketed"),
        Objective("extra", "Administración", 4, 0, 3, rule="manual")]
    c["objectives"] = [asdict(o) for o in objectives]
    return c


def test_sales_source_periods_scope_and_no_charge_exclusions(service):
    c = compact_campaign(service)
    c["scopes"]["cadelsa"].append("igsa")  # An alias selected twice must not double count.
    add_sales(service,
        ("igsa-old","2025-01","p",100,10,400,"ireks"),
        ("igsa","2025-12","p",200,20,800,"ireks"),
        ("igsa","2026-01","p",110,5,500,"ireks"),
        ("hermanos","2026-01","launch",10,30,40,"ireks"),
        ("igsa","2026-08","p",500,0,500,"ireks"),
        ("other","2026-12","p",10000,0,10000,"ireks"),
        ("igsa","2026-12","p",10000,0,10000,"igsa"))
    assert service.latest_month(c) == 8
    s = service.calculate(c, 1)
    rows = {r.objective.key:r for r in s["results"]}
    assert rows["total"].previous_year == 330
    assert rows["total"].annual_target == Decimal("339.90")
    assert rows["total"].previous == 110
    assert rows["total"].current == 155
    assert rows["familia"].current == 155
    assert rows["sc"].previous == 10
    assert rows["sc"].current == 5
    assert rows["sc"].points == 2
    assert rows["eur"].current == 540
    assert rows["pending"].current is None and rows["pending"].points == 0
    assert rows["ten"].current is None and rows["ten"].points == 0
    assert s["base_max"] == 29  # Pending and unmarketed maxima are retained.
    assert s["provisional"]
    assert not service.config_path.exists()


def test_missing_comparison_does_not_award_full_points(service):
    add_sales(service, ("igsa","2026-01","p",100,0,100,"ireks"))
    c = compact_campaign(service)
    r = service.calculate(c, 1)
    total = next(o for o in r["results"] if o.objective.key == "total")
    assert total.points is None
    assert total.status == "Sin datos comparables"


def test_configuration_year_isolation_and_validation(service):
    c = compact_campaign(service)
    c["manual"] = {"7": {"extra": 2}}
    service.save(c)
    assert service.campaign(2026) == c
    later = service.copy_campaign(2026, 2027)
    assert later["manual"] == {}
    later["objectives"][2]["members"] = ["p"]
    service.save(later)
    assert service.campaign(2026)["objectives"][2]["members"] == ["launch"]
    assert service.campaign(2027)["objectives"][2]["members"] == ["p"]
    original = service.config_path.read_bytes()
    later["objectives"][0]["maximum"] = -1
    with pytest.raises(ValueError):
        service.save(later)
    assert service.config_path.read_bytes() == original
    with pytest.raises(ValueError):
        service.copy_campaign(2026, 2027)


@pytest.mark.parametrize("current,expected,eligible", [(104,0,False), (105,3,True), (110,0,False)])
def test_manual_extra_eligibility(service, current, expected, eligible):
    c = compact_campaign(service)
    c["objectives"] = [asdict(Objective("base", "Base", 0, 10, 100)), asdict(Objective("extra", "Extra", 4, 0, 3, rule="manual"))]
    c["manual"] = {"1":{"extra":3}}
    add_sales(service, ("igsa","2025-01","p",100,0,0,"ireks"),("igsa","2026-01","p",current,0,0,"ireks"))
    s = service.calculate(c,1)
    assert s["extra"] == expected
    assert s["extra_eligible"] == eligible


def test_exports_preserve_numbers_and_pending_states(service, tmp_path):
    from app.services.objectives_export import export_objectives
    from openpyxl import load_workbook
    import fitz
    c = compact_campaign(service)
    add_sales(service,("igsa","2025-01","p",100,0,0,"ireks"),("igsa","2026-01","p",110,0,0,"ireks"))
    snapshot = service.calculate(c,1)
    xlsx = tmp_path / "objectives.xlsx"
    export_objectives(xlsx,snapshot)
    ws = load_workbook(xlsx).active
    assert ws.cell(5,5).value == 110
    assert ws.cell(10,5).value is None
    assert ws.cell(10,10).value == "Pendiente"
    assert ws.cell(5,5).data_type == "n"
    pdf = tmp_path / "objectives.pdf"
    export_objectives(pdf,snapshot)
    doc = fitz.open(pdf)
    text = "".join(page.get_text() for page in doc)
    assert "Pendiente" in text and "No comercializado" in text
    doc.close()


def test_page_defaults_to_latest_month_and_updates_saved_target(service):
    from PySide6.QtWidgets import QApplication
    from app.ui.widgets.objectives_page import ObjectivesPage
    app = QApplication.instance() or QApplication([])
    c = compact_campaign(service)
    service.save(c)
    add_sales(service,("igsa","2025-01","p",100,0,0,"ireks"),("igsa","2026-08","p",110,0,0,"ireks"))
    page = ObjectivesPage(service=service)
    page.reload()
    assert page.month.currentData() == 8
    assert page.detail.rowCount() == len(c["objectives"])
    changed = deepcopy(page.campaign)
    changed["objectives"][0]["target"] = 20
    assert page._save(changed)
    assert page.snapshot["results"][0].annual_target == 120
    page.detail.selectRow(0)
    assert "Desviación" in page.rule_note.text()
    page.filter.setCurrentIndex(page.filter.findData(1))
    assert page.detail.rowCount() == 2
    page.close()


def test_missing_family_is_not_zero_sales(service):
    c = compact_campaign(service)
    c["objectives"][1]["members"] = ["nonexistent"]
    snapshot = service.calculate(c,7)
    row = snapshot["results"][1]
    assert row.current is None and row.points is None
    assert snapshot["provisional"]


def test_no_charge_exclusions_apply_to_both_years(service):
    c = compact_campaign(service)
    add_sales(service,("igsa","2025-01","p",100,10,0,"ireks"),
        ("igsa","2025-01","launch",100,100,0,"ireks"),
        ("igsa","2026-01","p",100,12,0,"ireks"),
        ("igsa","2026-01","launch",100,0,0,"ireks"))
    sc = next(r for r in service.calculate(c,1)["results"] if r.objective.key == "sc")
    assert sc.previous == 10 and sc.current == 12
    assert sc.points == 0


def test_fixed_target_with_previous_sales_uses_derived_growth():
    objective = Objective("x", "Lanzamiento", 1, 1500, 4, target_type="kg")
    assert score(objective, 100, 125, 1000, 7) == 2
    assert score(objective, 100, 150, 1000, 7) == 4


def test_default_33d_includes_dreidoppel_and_gelatop(service):
    catalog = service.catalog()
    catalog["manufacturer"] = {"d": "DREIDOPPEL", "g": "GELATOP", "i": "IREKS"}
    campaign = service.default_campaign(catalog)
    objectives = {o["key"]: o for o in campaign["objectives"]}
    assert set(objectives["dreidoppel"]["members"]) == {"d", "g"}
    assert objectives["dreidoppel"]["unit"] == "kg"
    assert objectives["gelatop"]["members"] == ["g"]
    assert objectives["gelatop"]["unit"] == "eur"
