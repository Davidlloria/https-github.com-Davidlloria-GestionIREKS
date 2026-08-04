from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from app.core.database import engine


@dataclass
class CustomerImportPreviewItem:
    row_number: int
    action: str
    cliente_id: str = ""
    cliente_codigo_actual: int | None = None
    cliente_codigo_propuesto: int | None = None
    cliente_codigo_excel: int | None = None
    cliente_codigo_distribuidor: str = ""
    nombre_comercial_actual: str = ""
    nombre_comercial_propuesto: str = ""
    messages: list[str] = field(default_factory=list)


@dataclass
class CustomerImportPreviewResult:
    source_path: str
    total_rows: int
    items: list[CustomerImportPreviewItem]

    @property
    def updates(self) -> int:
        return sum(1 for item in self.items if item.action == "update")

    @property
    def creates(self) -> int:
        return sum(1 for item in self.items if item.action == "create")

    @property
    def skips(self) -> int:
        return sum(1 for item in self.items if item.action == "skip")

    @property
    def errors(self) -> int:
        return sum(1 for item in self.items if item.action == "error")

    def summary_lines(self) -> list[str]:
        return [
            f"Archivo: {self.source_path}",
            f"Filas utiles leidas: {self.total_rows}",
            f"Clientes existentes a actualizar: {self.updates}",
            f"Clientes nuevos a crear: {self.creates}",
            f"Filas omitidas: {self.skips}",
            f"Incidencias bloqueantes: {self.errors}",
        ]


@dataclass
class CustomerImportApplyResult:
    preview: CustomerImportPreviewResult
    updated: int
    created: int


class CustomerNewClientsImportPreviewService:
    def preview(self, source: str | Path) -> CustomerImportPreviewResult:
        path = Path(source)
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("El archivo debe ser .xlsx o .xlsm.")
        if not path.exists():
            raise FileNotFoundError(path)

        rows = self._read_rows(path)
        with engine.begin() as conn:
            customers_by_id = {
                str(row[0] or ""): {
                    "cliente_id": str(row[0] or ""),
                    "cliente_codigo": int(row[1] or 0),
                    "cliente_codigo_distribuidor": self._clean_text(row[2]),
                    "cliente_nombre_comercial": str(row[3] or ""),
                }
                for row in conn.exec_driver_sql(
                    """
                    SELECT cliente_id, cliente_codigo, cliente_codigo_distribuidor, cliente_nombre_comercial
                    FROM clientes
                    """
                ).fetchall()
            }
            used_codes = {
                int(row[0])
                for row in conn.exec_driver_sql(
                    "SELECT cliente_codigo FROM clientes WHERE cliente_codigo IS NOT NULL"
                ).fetchall()
                if self._to_int(row[0]) is not None
            }

        next_code = (max(used_codes) + 1) if used_codes else 1
        seen_distributor_codes: dict[str, int] = {}
        items: list[CustomerImportPreviewItem] = []

        for row in rows:
            messages: list[str] = []
            distributor_code_raw = row["cliente_codigo_distribuidor_raw"]
            distributor_code = row["cliente_codigo_distribuidor"]
            name = row["nombre_comercial"]
            if distributor_code_raw and not distributor_code:
                messages.append(
                    f"Codigo de cliente de distribuidor no valido en columna C: {distributor_code_raw}."
                )
            if distributor_code:
                previous_row = seen_distributor_codes.get(distributor_code)
                if previous_row:
                    messages.append(
                        f"Codigo distribuidor duplicado en Excel; ya aparece en fila {previous_row}."
                    )
                else:
                    seen_distributor_codes[distributor_code] = row["row_number"]

            if not distributor_code_raw:
                items.append(
                    CustomerImportPreviewItem(
                        row_number=row["row_number"],
                        action="skip",
                        cliente_id=row["cliente_id"],
                        cliente_codigo_excel=row["cliente_codigo"],
                        cliente_codigo_distribuidor="",
                        nombre_comercial_propuesto=name,
                        messages=["Sin codigo de cliente de distribuidor en columna C; no hay dato que aplicar."],
                    )
                )
                continue
            if not name:
                messages.append("Falta nombre comercial en columna D.")

            customer_id = row["cliente_id"]
            excel_code = row["cliente_codigo"]
            if customer_id:
                existing = customers_by_id.get(customer_id)
                if existing is None:
                    items.append(
                        CustomerImportPreviewItem(
                            row_number=row["row_number"],
                            action="error",
                            cliente_id=customer_id,
                            cliente_codigo_excel=excel_code,
                            cliente_codigo_distribuidor=distributor_code,
                            nombre_comercial_propuesto=name,
                            messages=[*messages, "UUID no encontrado en la tabla clientes."],
                        )
                    )
                    continue
                current_code = int(existing["cliente_codigo"] or 0)
                if excel_code is not None and excel_code != current_code:
                    messages.append(
                        f"Codigo interno no coincide: Excel={excel_code}, DB={current_code}."
                    )
                action = "error" if messages else "update"
                items.append(
                    CustomerImportPreviewItem(
                        row_number=row["row_number"],
                        action=action,
                        cliente_id=customer_id,
                        cliente_codigo_actual=current_code,
                        cliente_codigo_propuesto=current_code,
                        cliente_codigo_excel=excel_code,
                        cliente_codigo_distribuidor=distributor_code,
                        nombre_comercial_actual=str(existing["cliente_nombre_comercial"] or ""),
                        nombre_comercial_propuesto=name,
                        messages=messages,
                    )
                )
                continue

            if excel_code is not None:
                messages.append("Cliente nuevo con columna B informada; se esperaba codigo interno vacio.")
            if messages:
                items.append(
                    CustomerImportPreviewItem(
                        row_number=row["row_number"],
                        action="error",
                        cliente_codigo_excel=excel_code,
                        cliente_codigo_distribuidor=distributor_code,
                        nombre_comercial_propuesto=name,
                        messages=messages,
                    )
                )
                continue

            while next_code in used_codes:
                next_code += 1
            proposed_code = next_code
            used_codes.add(proposed_code)
            next_code += 1
            items.append(
                CustomerImportPreviewItem(
                    row_number=row["row_number"],
                    action="create",
                    cliente_codigo_propuesto=proposed_code,
                    cliente_codigo_distribuidor=distributor_code,
                    nombre_comercial_propuesto=name,
                )
            )

        return CustomerImportPreviewResult(str(path), len(rows), items)

    def _read_rows(self, path: Path) -> list[dict[str, Any]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows: list[dict[str, Any]] = []
        try:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = list(row[:4])
                while len(values) < 4:
                    values.append(None)
                if not any(self._clean_text(value) for value in values):
                    continue
                if self._looks_like_header(values):
                    continue
                rows.append(
                    {
                        "row_number": row_number,
                        "cliente_id": self._clean_text(values[0]),
                        "cliente_codigo": self._to_int(values[1]),
                        "cliente_codigo_distribuidor_raw": self._clean_text(values[2]),
                        "cliente_codigo_distribuidor": self._normalize_distributor_code(values[2]),
                        "nombre_comercial": self._clean_text(values[3]),
                    }
                )
        finally:
            workbook.close()
        return rows

    def _looks_like_header(self, values: list[Any]) -> bool:
        text = " ".join(self._clean_text(value).lower() for value in values)
        return (
            text.startswith("listado ")
            or "id cliente" in text
            or "uuid" in text
            or ("codigo" in text and "nombre" in text)
            or ("cod." in text and "nombre" in text)
        )

    def _clean_text(self, value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            return text[:-2]
        return text

    def _normalize_distributor_code(self, value: Any) -> str:
        text = self._clean_text(value)
        if not text:
            return ""
        if re.fullmatch(r"\d+", text):
            return text
        if re.fullmatch(r"\d+\s*-\s*\d+", text):
            return re.sub(r"\s*-\s*", "-", text)
        if re.fullmatch(r"\d+\s+\d+", text):
            return re.sub(r"\s+", "-", text)
        return ""

    def _to_int(self, value: Any) -> int | None:
        text = self._clean_text(value)
        if not text:
            return None
        try:
            return int(float(text.replace(",", ".")))
        except Exception:
            return None

    def apply(self, source: str | Path) -> CustomerImportApplyResult:
        preview = self.preview(source)
        if preview.errors:
            raise ValueError(
                f"No se puede aplicar la importacion: quedan {preview.errors} incidencias bloqueantes."
            )

        updated = 0
        created = 0
        with engine.begin() as conn:
            for item in preview.items:
                if item.action == "update":
                    conn.exec_driver_sql(
                        """
                        UPDATE clientes
                        SET cliente_codigo_distribuidor = ?, cliente_nombre_comercial = ?
                        WHERE cliente_id = ?
                        """,
                        (
                            item.cliente_codigo_distribuidor,
                            item.nombre_comercial_propuesto,
                            item.cliente_id,
                        ),
                    )
                    updated += 1
                elif item.action == "create":
                    conn.exec_driver_sql(
                        """
                        INSERT INTO clientes (
                            cliente_id, cliente_codigo, cliente_codigo_distribuidor,
                            cliente_nombre_comercial, cliente_nombre_fiscal,
                            cliente_nombre_interno, cliente_abreviatura, cliente_cif,
                            cliente_telefono, cliente_email, cliente_direccion,
                            cliente_direccion_cp, cliente_direccion_localidad_id,
                            cliente_direccion_municipio_id, cliente_direccion_provincia_id,
                            cliente_direccion_isla_id, cliente_tipo, cliente_actividad,
                            cliente_prospeccion, distribuidor_id, distribuidor_comercial_id,
                            activo
                        )
                        VALUES (?, ?, ?, ?, ?, '', '', '', '', '', '', '', '', '', '', '', '', '', 0, '', '', 1)
                        """,
                        (
                            str(uuid4()),
                            item.cliente_codigo_propuesto,
                            item.cliente_codigo_distribuidor,
                            item.nombre_comercial_propuesto,
                            item.nombre_comercial_propuesto,
                        ),
                    )
                    created += 1
        return CustomerImportApplyResult(preview=preview, updated=updated, created=created)
