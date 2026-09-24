"""Read-only historical comparison; never creates inventory adjustments."""
from collections import Counter, defaultdict
from datetime import date, datetime
from math import isfinite

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.core.database import engine
from app.models import AlmacenMovimiento, IngredienteIreks


def code_key(value):
    text = str(value if value is not None else "").strip().upper()
    return str(int(text)) if text.isdigit() else text


def lot_key(value):
    return str(value if value is not None else "").strip().upper()


class InventoryReviewService:
    def __init__(self, db_engine=None):
        self.engine = db_engine if db_engine is not None else engine

    def compare(self, path, warehouse, cutoff):
        if not warehouse:
            raise ValueError("Selecciona un almacén concreto antes de comparar.")
        book = load_workbook(path, data_only=False)
        try:
            if "INVENTARIO" not in book.sheetnames:
                raise ValueError("Falta la hoja INVENTARIO con el conteo físico definitivo.")
            sheet = book["INVENTARIO"]
            headers = [sheet.cell(3, c).value for c in range(1, 15)]
            if headers[1:5] != ["Artículo", "Descripción", "Lote", "Fecha cons.pref."] or headers[7] != "Conteo Envases":
                raise ValueError("La estructura de INVENTARIO no coincide con el formato esperado.")
            source = [(i, list(values)) for i, values in enumerate(sheet.values, 1)
                      if i > 3 and any(v is not None for v in values[:8])]
        finally:
            book.close()
        with Session(self.engine) as session:
            products = list(session.exec(select(IngredienteIreks)))
            moves = list(session.exec(select(AlmacenMovimiento).where(
                AlmacenMovimiento.almacen_id == warehouse,
                AlmacenMovimiento.fecha_pedido <= cutoff)))
        aliases = defaultdict(dict)
        by_id = {p.articulo_id: p for p in products}
        for product in products:
            for ref in (product.articulo_referencia, product.articulo_referencia_corta):
                if ref:
                    aliases[code_key(ref)][product.articulo_id] = product
        balances = defaultdict(float)
        for move in moves:
            balances[move.articulo_id, lot_key(move.articulo_lote)] += move.cantidad
        duplicates = Counter((code_key(r[1]), lot_key(r[3])) for _, r in source)
        result, represented = [], set()
        for number, values in source:
            ref, name, lot, expiry, count = values[1], values[2], values[3], values[4], values[7]
            candidates = aliases.get(code_key(ref), {})
            product = next(iter(candidates.values())) if len(candidates) == 1 else None
            issues = []
            if product is None:
                issues.append("Producto no identificado" if not candidates else "Código ambiguo")
            if not lot_key(lot):
                issues.append("Lote pendiente")
            if duplicates[code_key(ref), lot_key(lot)] > 1:
                issues.append("Producto y lote duplicados")
            if count is None or count == "":
                issues.append("Conteo pendiente; excluido")
                count = None
            elif isinstance(count, bool) or not isinstance(count, (int, float)) or not isfinite(count) or count < 0:
                issues.append("Conteo no válido; excluido")
                count = None
            key = (product.articulo_id, lot_key(lot)) if product else None
            if key:
                represented.add(key)
            theoretical = balances.get(key, 0) if key else None
            weight = float(product.articulo_envase_peso_total or product.articulo_envase_peso or 0) if product else 0
            delta = count - theoretical if not issues else None
            if theoretical is not None and theoretical < -1e-6:
                issues.append("Saldo teórico negativo; revisar movimientos")
            if weight <= 0:
                issues.append("Peso pendiente; kg no calculados")
            expired = isinstance(expiry, (date, datetime)) and (expiry.date() if isinstance(expiry, datetime) else expiry) < cutoff
            notes = str(values[13] or "")
            if expired or "caducad" in (str(expiry) + notes).lower():
                issues.append("Caducado: incluido en físico; revisar destrucción")
            result.append(dict(row=number, code=str(ref or ""), name=str(name or ""), lot=lot_key(lot),
                               expiry=str(expiry or "")[:10], theoretical=theoretical, count=count,
                               delta=delta, kg=delta * weight if delta is not None and weight > 0 else None,
                               status="; ".join(issues) or ("Coincide" if abs(delta) < 1e-6 else "Diferencia por revisar"), notes=notes))
        for (product_id, lot), quantity in sorted(balances.items()):
            if (product_id, lot) in represented or abs(quantity) < 1e-6:
                continue
            product = by_id.get(product_id)
            result.append(dict(row=None, code=product.articulo_referencia if product else product_id,
                               name=product.articulo_descripcion if product else "Producto sin ficha", lot=lot,
                               expiry="", theoretical=quantity, count=None, delta=None, kg=None,
                               status="No figura en el inventario; pendiente, no se supone cero", notes=""))
        return result
