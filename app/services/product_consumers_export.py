"""Exports of the product consumers comparison, in the visible sort order."""
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


TITLE = "Clientes que compran el producto"
NAVY = "173653"


def format_value(value, column):
    prefix = "+" if column >= 6 and value > 0 else ""
    text = f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return prefix + text + (" €" if column in (3, 5, 7) else "")


def export_excel(path, subtitle, headers, rows, totals):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes del producto"
    for row_number, text in ((1, TITLE), (2, subtitle)):
        ws.cell(row_number, 1, text)
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=8)
    ws.cell(1, 1).font = Font(size=17, bold=True, color=NAVY)
    ws.append(headers)
    for values in rows:
        ws.append(values)
        # Customer text is literal, including names starting with '='.
        for col in (1, 2):
            ws.cell(ws.max_row, col).data_type = "s"
    end = ws.max_row
    ws.append(["Totales generales", "", *totals])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
    for row in ws.iter_rows(min_row=3):
        is_band = row[0].row in (3, ws.max_row)
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY if is_band else ("F3F6F9" if cell.row % 2 == 0 else "FFFFFF"))
            color = "FFFFFF" if is_band else NAVY
            if cell.column >= 7 and isinstance(cell.value, (float, int)) and cell.value:
                color = ("76E2B6" if cell.value > 0 else "FF8585") if is_band else ("07804B" if cell.value > 0 else "C32939")
            cell.font = Font(name="Calibri", size=11, bold=is_band, color=color)
            cell.alignment = Alignment(horizontal="right" if cell.column >= 3 else "left", vertical="center", wrap_text=cell.column == 2)
            if cell.row > 3 and cell.column >= 3:
                unit = '" €"' if cell.column in (4, 6, 8) else ""
                cell.number_format = (f'+#,##0.00{unit};-#,##0.00{unit};0.00{unit}' if cell.column >= 7 else f'#,##0.00{unit}')
    for col, width in enumerate((11, 48, 18, 21, 18, 21, 18, 21), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:H{end}"
    ws.print_title_rows = "1:3"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    wb.save(path)


def export_pdf(path, subtitle, headers, rows, totals):
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=22, rightMargin=22, topMargin=22, bottomMargin=28)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ConsumersTitle", parent=styles["Heading1"], textColor=colors.HexColor("#" + NAVY), fontSize=17)
    text = ParagraphStyle("ConsumersCell", parent=styles["Normal"], fontSize=7, leading=9)
    heading = ParagraphStyle("ConsumersHeader", parent=text, textColor=colors.white, fontName="Helvetica-Bold")
    centered_heading = ParagraphStyle("ConsumersNumericHeader", parent=heading, alignment=TA_CENTER)
    # Paragraphs wrap long customer names instead of drawing into numeric cells.
    previous_year = headers[2].split("·", 1)[0].strip()
    current_year = headers[4].split("·", 1)[0].strip()
    group_labels = [headers[0], headers[1], previous_year, "", current_year, "", "Diferencias", ""]
    data = [
        [Paragraph(escape(label), heading if col < 2 else centered_heading) for col, label in enumerate(group_labels)],
        ["", "", *[Paragraph(label, centered_heading) for label in ("Kilos", "€", "Kilos", "€", "Δ Kilos", "Δ €")]],
    ]
    for values in rows:
        data.append([Paragraph(escape(str(value)), text) if col < 2 else format_value(value, col) for col, value in enumerate(values)])
    data.append(["Totales generales", "", *[format_value(value, col) for col, value in enumerate(totals, 2)]])
    widths = [doc.width * factor for factor in (.055, .275, .105, .12, .105, .12, .10, .12)]
    table = Table(data, colWidths=widths, repeatRows=2, hAlign="LEFT")
    commands = [
        ("SPAN", (0, 0), (0, 1)),
        ("SPAN", (1, 0), (1, 1)),
        ("SPAN", (2, 0), (3, 0)),
        ("SPAN", (4, 0), (5, 0)),
        ("SPAN", (6, 0), (7, 0)),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#" + NAVY)),
        ("ROWBACKGROUNDS", (0, 2), (-1, -2), [colors.white, colors.HexColor("#F3F6F9")]),
        ("TEXTCOLOR", (0, 2), (-1, -2), colors.HexColor("#" + NAVY)),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("SPAN", (0, -1), (1, -1)),
        ("ALIGN", (2, 2), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (2, 0), (-1, 0), 0.5, colors.HexColor("#70879A")),
    ]
    for col in range(1, 8):
        # Pair dividers start below the merged year labels; the total label spans
        # the first two columns and must not have a line through its text.
        commands.append(("LINEBEFORE", (col, 1 if col in (3, 5, 7) else 0), (col, -2 if col == 1 else -1), 0.4, colors.HexColor("#A9B9C8")))
    for row_idx, values in enumerate([*rows, ["", "", *totals]], 2):
        for col in (6, 7):
            if values[col]:
                palette = ("#76E2B6", "#FF8585") if row_idx == len(data) - 1 else ("#07804B", "#C32939")
                commands.append(("TEXTCOLOR", (col, row_idx), (col, row_idx), colors.HexColor(palette[values[col] < 0])))
    table.setStyle(TableStyle(commands))

    def page_number(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#60758A"))
        canvas.drawRightString(landscape(A4)[0] - 22, 14, f"Página {document.page}")

    doc.build([Paragraph(TITLE, title), Paragraph(escape(subtitle), styles["Normal"]), Spacer(1, 14), table], onFirstPage=page_number, onLaterPages=page_number)
