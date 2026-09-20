from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.objectives_service import MONTHS, fmt


def export_objectives(path, snapshot):
    year, month = snapshot["year"], snapshot["month"]
    title = f"Objetivos Canarias {year} · acumulado hasta {MONTHS[month-1]}"
    subtitle = f'Ventas IREKS · kg vendidos + sin cargo · Puntos: {fmt(snapshot["total"])} · Base: {fmt(snapshot["base"])} / {fmt(snapshot["base_max"])} · Extra: {fmt(snapshot["extra"])}'
    headers = ["Objetivo", "Ud.", "Meta anual", f"{year-1}", f"{year}", "Δ cantidad", "Δ %", "Puntos", "Máximo", "Estado"]
    rows = []
    for r in snapshot["results"]:
        unit = "—" if r.objective.rule == "manual" else ("€" if r.objective.unit == "eur" else "kg")
        rows.append([r.objective.name, unit, r.annual_target, r.previous, r.current, r.delta, r.growth, r.points, r.objective.maximum, r.status])
    notes = "Las cantidades de los apartados se solapan y no se suman. Extra aplicable desde el 45 % y por debajo del 100 % de la base."
    if snapshot["provisional"]:
        notes += " Resultado provisional. " + " · ".join(snapshot["issues"])
    if Path(path).suffix.lower() == ".xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Objetivos"
        for text in (title, subtitle, notes):
            ws.append([text])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=10)
            ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[3].height = 42
        ws.append(headers)
        for values in rows:
            ws.append([float(v) if v is not None and col in range(2, 9) else v for col, v in enumerate(values)])
            for col in (1, 2, 10):
                ws.cell(ws.max_row, col).data_type = "s"
        ws.freeze_panes = "C5"
        ws.auto_filter.ref = f"A4:J{ws.max_row}"
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.font = Font(name="Calibri", size=11, bold=cell.row == 4, color="FFFFFF" if cell.row == 4 else "173653")
                cell.fill = PatternFill("solid", fgColor="173653" if cell.row == 4 else ("F3F6F9" if cell.row % 2 else "FFFFFF"))
                cell.alignment = Alignment(horizontal="center" if cell.row == 4 else ("right" if 3 <= cell.column <= 9 else "left"))
                if cell.row > 4 and 3 <= cell.column <= 9:
                    cell.number_format = '#,##0.00" %"' if cell.column == 7 else '#,##0.00'
        for col, width in enumerate((52, 9, 19, 19, 19, 19, 12, 12, 12, 25), 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.print_title_rows = "1:4"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        wb.save(path)
        return
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=22, rightMargin=22, topMargin=22, bottomMargin=28)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ObjectiveCell", parent=styles["Normal"], fontSize=7, leading=9)
    header = ParagraphStyle("ObjectiveHeader", parent=body, alignment=TA_CENTER, textColor=colors.white, fontName="Helvetica-Bold")
    data = [[Paragraph(escape(h), header) for h in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(v or "")), body) if col in (0, 1, 9) else fmt(v) for col, v in enumerate(row)])
    widths = [doc.width*x for x in (.25, .04, .105, .105, .105, .105, .06, .055, .055, .12)]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173653")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F9")]),
        ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#B9C8D6")),
        ("ALIGN", (2, 1), (8, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(landscape(A4)[0]-22, 14, f"Página {document.page}")
    doc.build([Paragraph(escape(title), styles["Heading1"]), Paragraph(escape(subtitle), body), Spacer(1, 10), table, Spacer(1, 10), Paragraph(escape(notes), body)], onFirstPage=footer, onLaterPages=footer)
