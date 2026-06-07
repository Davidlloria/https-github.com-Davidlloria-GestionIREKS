from __future__ import annotations

from datetime import date
import html
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from PySide6.QtCore import QDate, QSize, QTimer, Qt, QMarginsF
from PySide6.QtGui import QIcon, QPageLayout, QPageSize, QPagedPaintDevice, QPainter, QPdfWriter, QTextDocument, QPixmap
from PySide6.QtPdf import QPdfDocument
from PySide6.QtPdfWidgets import QPdfView
from PySide6.QtPrintSupport import QPrintDialog, QPrinter
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDateEdit,
    QDialog,
    QFileDialog,
    QFrame,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPlainTextEdit,
    QPushButton,
    QDialogButtonBox,
    QStackedWidget,
    QSplitter,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from app.services.course_document_generation_flow_service import CourseDocumentGenerationFlowService
from app.services.certificate_service import CertificateService
from app.services.course_service import CourseService
from app.services.signature_sheet_service import SignatureSheetService
from app.ui.widgets.entity_dialog import EntityDialog

PENCIL_ICON_PATH = Path(__file__).resolve().parents[3] / "assets" / "icons" / "pencil_white.svg"


class AttendeePickerDialog(QDialog):
    def __init__(self, service: CourseService, parent=None) -> None:
        super().__init__(parent)
        self.service = service
        self.rows = []
        self._search_timer = QTimer(self)
        self._search_timer.setSingleShot(True)
        self._search_timer.timeout.connect(self._reload)
        self.setWindowTitle("Seleccionar asistente")
        self.resize(920, 620)
        self._build_ui()
        self._reload()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        top = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Filtrar por nombre o empresa...")
        self.search_input.textChanged.connect(lambda *_: self._search_timer.start(250))
        top.addWidget(self.search_input, 1)
        layout.addLayout(top)

        self.table = QTableWidget(0, 2)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table.setHorizontalHeaderLabels(["Nombre", "Empresa"])
        self.table.setSortingEnabled(True)
        self.table.doubleClicked.connect(self.accept)
        layout.addWidget(self.table, 1)

        actions = QHBoxLayout()
        use_btn = QPushButton("Seleccionar")
        use_btn.setProperty("btnRole", "success")
        cancel_btn = QPushButton("Cancelar")
        cancel_btn.setProperty("btnRole", "secondary")
        use_btn.clicked.connect(self.accept)
        cancel_btn.clicked.connect(self.reject)
        actions.addWidget(use_btn)
        actions.addWidget(cancel_btn)
        actions.addStretch(1)
        layout.addLayout(actions)

    def _reload(self) -> None:
        self.rows = self.service.list_contacts_for_picker(self.search_input.text().strip())
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(self.rows))
        for i, row in enumerate(self.rows):
            name_cell = QTableWidgetItem(row.nombre_completo)
            name_cell.setData(Qt.ItemDataRole.UserRole, row.contacto_id)
            self.table.setItem(i, 0, name_cell)
            self.table.setItem(i, 1, QTableWidgetItem(row.empresa))
        self.table.setSortingEnabled(True)

    def selected(self):
        selected = self.table.selectionModel().selectedRows()
        if not selected:
            return None
        item = self.table.item(selected[0].row(), 0)
        if item is None:
            return None
        contacto_id = str(item.data(Qt.ItemDataRole.UserRole) or "")
        return next((r for r in self.rows if r.contacto_id == contacto_id), None)


class TechnicianPickerDialog(QDialog):
    def __init__(self, service: CourseService, parent=None) -> None:
        super().__init__(parent)
        self.service = service
        self.rows = []
        self._search_timer = QTimer(self)
        self._search_timer.setSingleShot(True)
        self._search_timer.timeout.connect(self._reload)
        self.setWindowTitle("Seleccionar tecnico")
        self.resize(920, 620)
        self._build_ui()
        self._reload()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        top = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Filtrar por nombre, movil, interno o email...")
        self.search_input.textChanged.connect(lambda *_: self._search_timer.start(250))
        top.addWidget(self.search_input, 1)
        layout.addLayout(top)

        self.table = QTableWidget(0, 4)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.table.setHorizontalHeaderLabels(["Tecnico", "Movil", "Interno", "Email"])
        self.table.setSortingEnabled(True)
        self.table.doubleClicked.connect(self.accept)
        layout.addWidget(self.table, 1)

        actions = QHBoxLayout()
        use_btn = QPushButton("Seleccionar")
        use_btn.setProperty("btnRole", "success")
        cancel_btn = QPushButton("Cancelar")
        cancel_btn.setProperty("btnRole", "secondary")
        use_btn.clicked.connect(self.accept)
        cancel_btn.clicked.connect(self.reject)
        actions.addWidget(use_btn)
        actions.addWidget(cancel_btn)
        actions.addStretch(1)
        layout.addLayout(actions)

    def _reload(self) -> None:
        self.rows = self.service.list_technicians_for_picker(self.search_input.text().strip())
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(self.rows))
        for i, row in enumerate(self.rows):
            name_cell = QTableWidgetItem(row.nombre_completo)
            name_cell.setData(Qt.ItemDataRole.UserRole, row.tecnico_id)
            self.table.setItem(i, 0, name_cell)
            self.table.setItem(i, 1, QTableWidgetItem(row.movil))
            self.table.setItem(i, 2, QTableWidgetItem(row.interno))
            self.table.setItem(i, 3, QTableWidgetItem(row.email))
        self.table.setSortingEnabled(True)

    def selected(self):
        selected = self.table.selectionModel().selectedRows()
        if not selected:
            return None
        item = self.table.item(selected[0].row(), 0)
        if item is None:
            return None
        tecnico_id = str(item.data(Qt.ItemDataRole.UserRole) or "")
        return next((r for r in self.rows if r.tecnico_id == tecnico_id), None)


class DocumentExpandedDialog(QDialog):
    def __init__(self, title: str, path: str, parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(1000, 760)
        layout = QVBoxLayout(self)
        view = CoursesPage.build_document_preview_widget(path, parent=self)
        layout.addWidget(view)


class ObservacionesDialog(QDialog):
    def __init__(self, title: str, label: str, initial: str = "", parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(760, 360)
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(label))
        self.input = QPlainTextEdit()
        self.input.setPlainText(initial)
        layout.addWidget(self.input, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def value(self) -> str:
        return self.input.toPlainText().strip()


class ConsentimientosDialog(QDialog):
    def __init__(self, parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Consentimientos")
        self.resize(520, 220)
        layout = QVBoxLayout(self)

        row_template = QHBoxLayout()
        row_template.addWidget(QLabel("Documento"))
        self.template_combo = QComboBox()
        self.template_combo.addItem("Imágenes", "imagenes")
        self.template_combo.addItem("Datos", "datos")
        row_template.addWidget(self.template_combo, 1)
        layout.addLayout(row_template)

        row_scope = QHBoxLayout()
        row_scope.addWidget(QLabel("Alcance"))
        self.scope_combo = QComboBox()
        self.scope_combo.addItem("Todos", "all")
        self.scope_combo.addItem("Solo confirmados", "confirmed")
        self.scope_combo.addItem("Hoja seleccionada", "selected")
        row_scope.addWidget(self.scope_combo, 1)
        layout.addLayout(row_scope)

        actions = QHBoxLayout()
        self.preview_btn = QPushButton("Previsualizar")
        self.preview_btn.setProperty("btnRole", "secondary")
        self.print_btn = QPushButton("Imprimir")
        self.print_btn.setProperty("btnRole", "secondary")
        self.close_btn = QPushButton("Cerrar")
        self.close_btn.setProperty("btnRole", "secondary")
        self.close_btn.clicked.connect(self.accept)
        actions.addWidget(self.preview_btn)
        actions.addWidget(self.print_btn)
        actions.addStretch(1)
        actions.addWidget(self.close_btn)
        layout.addLayout(actions)

    def selected_template(self) -> str:
        return str(self.template_combo.currentData() or "imagenes")

    def selected_scope(self) -> str:
        return str(self.scope_combo.currentData() or "all")


class CertificadosDialog(QDialog):
    def __init__(self, parent=None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Certificados")
        self.resize(520, 220)
        layout = QVBoxLayout(self)

        row_scope = QHBoxLayout()
        row_scope.addWidget(QLabel("Alcance"))
        self.scope_combo = QComboBox()
        self.scope_combo.addItem("Todos", "all")
        self.scope_combo.addItem("Solo confirmados", "confirmed")
        self.scope_combo.addItem("Hoja seleccionada", "selected")
        row_scope.addWidget(self.scope_combo, 1)
        layout.addLayout(row_scope)

        actions = QHBoxLayout()
        self.preview_btn = QPushButton("Previsualizar")
        self.preview_btn.setProperty("btnRole", "secondary")
        self.print_btn = QPushButton("Imprimir")
        self.print_btn.setProperty("btnRole", "secondary")
        self.close_btn = QPushButton("Cerrar")
        self.close_btn.setProperty("btnRole", "secondary")
        self.close_btn.clicked.connect(self.accept)
        actions.addWidget(self.preview_btn)
        actions.addWidget(self.print_btn)
        actions.addStretch(1)
        actions.addWidget(self.close_btn)
        layout.addLayout(actions)

    def selected_scope(self) -> str:
        return str(self.scope_combo.currentData() or "all")


class CoursesPage(QWidget):
    MONTHS = [
        (1, "Enero"),
        (2, "Febrero"),
        (3, "Marzo"),
        (4, "Abril"),
        (5, "Mayo"),
        (6, "Junio"),
        (7, "Julio"),
        (8, "Agosto"),
        (9, "Septiembre"),
        (10, "Octubre"),
        (11, "Noviembre"),
        (12, "Diciembre"),
    ]

    def __init__(self) -> None:
        super().__init__()
        self.service = CourseService()
        self.signature_service = SignatureSheetService()
        self.certificate_service = CertificateService()
        self.course_document_generation_service = CourseDocumentGenerationFlowService(
            signature_service=self.signature_service,
            certificate_service=self.certificate_service,
        )
        self.rows = []
        self.attendee_rows = []
        self.technician_rows = []
        self._is_loading_details = False
        self._loading_attendees = False
        self._attendee_sort_col = 1
        self._attendee_sort_order = Qt.SortOrder.AscendingOrder
        self._search_timer = QTimer(self)
        self._search_timer.setSingleShot(True)
        self._search_timer.timeout.connect(self.reload)
        self._autosave_timer = QTimer(self)
        self._autosave_timer.setSingleShot(True)
        self._autosave_timer.timeout.connect(self._autosave_selected_course)
        self._build_ui()
        self.reload()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        header = QLabel("Cursos")
        header.setProperty("role", "pageTitle")
        layout.addWidget(header)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        layout.addWidget(splitter, 1)

        left_panel = QWidget()
        left_panel.setObjectName("sidePanel")
        left_layout = QVBoxLayout(left_panel)
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Buscar por nombre de curso...")
        self.search_input.textChanged.connect(self._schedule_reload)
        left_layout.addWidget(self.search_input)

        filters_row = QHBoxLayout()
        self.year_filter = QComboBox()
        self.year_filter.addItem("Ano (todos)", None)
        for y in range(2025, date.today().year + 6):
            self.year_filter.addItem(str(y), y)
        self.year_filter.currentIndexChanged.connect(self._schedule_reload)
        self.month_start_filter = QComboBox()
        self.month_start_filter.addItem("Mes inicial", None)
        self.month_end_filter = QComboBox()
        self.month_end_filter.addItem("Mes final", None)
        for m, n in self.MONTHS:
            self.month_start_filter.addItem(n, m)
            self.month_end_filter.addItem(n, m)
        self.month_start_filter.currentIndexChanged.connect(self._schedule_reload)
        self.month_end_filter.currentIndexChanged.connect(self._schedule_reload)
        current_year_idx = self.year_filter.findData(date.today().year)
        if current_year_idx >= 0:
            self.year_filter.setCurrentIndex(current_year_idx)
        filters_row.addWidget(self.year_filter, 1)
        filters_row.addWidget(self.month_start_filter, 1)
        filters_row.addWidget(self.month_end_filter, 1)
        left_layout.addLayout(filters_row)

        self.table = QTableWidget(0, 2)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        table_header = self.table.horizontalHeader()
        table_header.setSectionsClickable(True)
        table_header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        table_header.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table.setHorizontalHeaderLabels(["Fecha", "Nombre"])
        self.table.setSortingEnabled(True)
        self.table.itemSelectionChanged.connect(self._show_selected_details)
        left_layout.addWidget(self.table, 1)
        splitter.addWidget(left_panel)

        right_panel = QWidget()
        right_layout = QVBoxLayout(right_panel)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_layout.setSpacing(6)
        ribbon = QFrame()
        ribbon.setObjectName("topRibbon")
        ribbon_layout = QHBoxLayout(ribbon)
        ribbon_layout.setContentsMargins(8, 6, 8, 6)
        new_btn = QPushButton("Nuevo")
        new_btn.setProperty("btnRole", "success")
        del_btn = QPushButton("Eliminar")
        del_btn.setProperty("btnRole", "danger")
        import_btn = QPushButton("Importar Excel/CSV")
        import_btn.setProperty("btnRole", "secondary")
        refresh_btn = QPushButton("Refrescar")
        refresh_btn.setProperty("btnRole", "secondary")
        new_btn.clicked.connect(self._new_course)
        del_btn.clicked.connect(self._delete_course)
        import_btn.clicked.connect(self._import_courses)
        refresh_btn.clicked.connect(self.reload)
        ribbon_layout.addWidget(new_btn)
        ribbon_layout.addWidget(del_btn)
        ribbon_layout.addWidget(import_btn)
        ribbon_layout.addWidget(refresh_btn)
        ribbon_layout.addStretch(1)
        right_layout.addWidget(ribbon)

        right_splitter = QSplitter(Qt.Orientation.Vertical)
        right_layout.addWidget(right_splitter, 1)
        detail_panel = QWidget()
        detail_panel.setObjectName("detailPanel")
        detail_layout = QVBoxLayout(detail_panel)
        detail_layout.setContentsMargins(14, 14, 14, 14)
        detail_title = QLabel("Detalle de curso")
        detail_title.setProperty("role", "sectionTitle")
        detail_layout.addWidget(detail_title)
        row_1 = QHBoxLayout()
        self.detail_curso_nombre = QLineEdit()
        self.detail_curso_fecha = QDateEdit()
        self.detail_curso_fecha.setCalendarPopup(True)
        self.detail_curso_fecha.setDisplayFormat("dd/MM/yyyy")
        row_1.addWidget(QLabel("Nombre"))
        row_1.addWidget(self.detail_curso_nombre, 2)
        row_1.addWidget(QLabel("Fecha"))
        row_1.addWidget(self.detail_curso_fecha, 1)
        detail_layout.addLayout(row_1)
        self.detail_curso_nombre.textEdited.connect(self._schedule_autosave)
        self.detail_curso_fecha.dateChanged.connect(self._schedule_autosave)
        right_splitter.addWidget(detail_panel)

        tabs_panel = QWidget()
        tabs_layout = QVBoxLayout(tabs_panel)
        tabs_layout.setContentsMargins(0, 0, 0, 0)
        self.tabs = QTabWidget()
        self.tabs.addTab(self._build_attendees_tab(), "Asistentes")
        self.tabs.addTab(self._build_technicians_tab(), "Tecnicos")
        self.tabs.addTab(self._build_documents_tab(), "Documentos")
        tabs_layout.addWidget(self.tabs)
        right_splitter.addWidget(tabs_panel)
        right_splitter.setStretchFactor(0, 4)
        right_splitter.setStretchFactor(1, 6)

        splitter.addWidget(right_panel)
        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 1)
        splitter.setChildrenCollapsible(False)
        splitter.setHandleWidth(0)
        splitter.handle(1).setEnabled(False)

    def _build_placeholder_tab(self, text: str) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        label = QLabel(text)
        label.setWordWrap(True)
        layout.addWidget(label, 1)
        return panel

    def _build_attendees_tab(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        actions = QHBoxLayout()
        for label, role, handler in [
            ("Añadir", "success", self._add_attendee),
            ("Eliminar", "danger", self._delete_attendee),
            ("Importar", "secondary", self._import_attendees),
            ("Exportar Excel", "secondary", self._export_attendees_excel),
            ("Exportar PDF", "secondary", self._export_attendees_pdf),
            ("Consentimientos", "secondary", self._open_consentimientos_manager),
            ("Certificado", "secondary", self._open_certificados_manager),
        ]:
            btn = QPushButton(label)
            btn.setProperty("btnRole", role)
            btn.clicked.connect(handler)
            actions.addWidget(btn)
        actions.addStretch(1)
        self.attendees_counter_box = QWidget()
        self.attendees_counter_box.setStyleSheet(
            "QWidget { background: transparent; border: 1px solid #D7DEE8; border-radius: 6px; }"
        )
        self.attendees_counter_box.setFixedHeight(30)
        self.attendees_counter_box.setFixedWidth(100)
        counter_layout = QHBoxLayout(self.attendees_counter_box)
        counter_layout.setContentsMargins(10, 0, 10, 0)
        self.attendees_counter_label = QLabel("0/0")
        self.attendees_counter_label.setStyleSheet("QLabel { border: none; background: transparent; }")
        self.attendees_counter_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        counter_layout.addWidget(self.attendees_counter_label, 1)
        actions.addWidget(self.attendees_counter_box)
        layout.addLayout(actions)
        self.attendees_table = QTableWidget(0, 5)
        self.attendees_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.attendees_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.attendees_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.attendees_table.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        self.attendees_table.verticalHeader().setVisible(False)
        self.attendees_table.verticalHeader().setDefaultSectionSize(35)
        self.attendees_table.verticalHeader().setMinimumSectionSize(35)
        self.attendees_table.setStyleSheet("QTableWidget::item:focus { border: none; }")
        header = self.attendees_table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.Interactive)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(4, QHeaderView.ResizeMode.Fixed)
        self.attendees_table.setHorizontalHeaderLabels(["Asiste", "Asistente", "NIF", "Empresa", ""])
        self.attendees_table.setColumnWidth(1, 300)
        self.attendees_table.setColumnWidth(4, 52)
        self.attendees_table.setSortingEnabled(False)
        header.sectionClicked.connect(self._sort_attendees_by_column)
        header.setSortIndicatorShown(True)
        header.setSortIndicator(self._attendee_sort_col, self._attendee_sort_order)
        self.attendees_table.itemChanged.connect(self._on_attendee_item_changed)
        self.attendees_table.cellDoubleClicked.connect(self._open_attendee_contact)
        layout.addWidget(self.attendees_table, 1)
        return panel

    def _build_technicians_tab(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        actions = QHBoxLayout()
        for label, role, handler in [
            ("Añadir", "success", self._add_technician),
            ("Eliminar", "danger", self._delete_technician),
        ]:
            btn = QPushButton(label)
            btn.setProperty("btnRole", role)
            btn.clicked.connect(handler)
            actions.addWidget(btn)
        actions.addStretch(1)
        layout.addLayout(actions)

        self.technicians_table = QTableWidget(0, 4)
        self.technicians_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.technicians_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.technicians_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.technicians_table.verticalHeader().setVisible(False)
        header = self.technicians_table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.technicians_table.setHorizontalHeaderLabels(["Tecnico", "Movil", "Interno", "Email"])
        self.technicians_table.setSortingEnabled(True)
        layout.addWidget(self.technicians_table, 1)
        return panel

    def _build_documents_tab(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        self.doc_paths = {"portada": "", "invitacion": "", "recetario": ""}
        self.doc_previews = {}
        self.doc_tabs = QTabWidget()
        doc_defs = [("portada", "Portada"), ("invitacion", "Invitación"), ("recetario", "Recetario")]
        for field, label in doc_defs:
            tab = QWidget()
            tab_layout = QVBoxLayout(tab)
            actions = QHBoxLayout()
            attach_btn = QPushButton("Adjuntar")
            attach_btn.setProperty("btnRole", "secondary")
            attach_btn.clicked.connect(lambda _checked=False, f=field: self._attach_document(f))
            delete_btn = QPushButton("Eliminar")
            delete_btn.setProperty("btnRole", "danger")
            delete_btn.clicked.connect(lambda _checked=False, f=field: self._delete_document(f))
            print_btn = QPushButton("Imprimir")
            print_btn.setProperty("btnRole", "secondary")
            print_btn.clicked.connect(lambda _checked=False, f=field: self._print_document(f))
            expand_btn = QPushButton("Ampliar")
            expand_btn.setProperty("btnRole", "secondary")
            expand_btn.clicked.connect(lambda _checked=False, f=field: self._expand_document(f))
            actions.addWidget(attach_btn)
            actions.addWidget(delete_btn)
            actions.addWidget(print_btn)
            actions.addWidget(expand_btn)
            actions.addStretch(1)
            tab_layout.addLayout(actions)

            preview = self.build_document_preview_widget("", parent=self)
            tab_layout.addWidget(preview, 1)
            self.doc_previews[field] = preview
            self.doc_tabs.addTab(tab, label)
        layout.addWidget(self.doc_tabs, 1)
        return panel

    def _selected_course(self):
        selected = self.table.selectionModel().selectedRows()
        if not selected:
            return None
        item = self.table.item(selected[0].row(), 0)
        if item is None:
            return None
        course_id = item.data(Qt.ItemDataRole.UserRole)
        return next((r for r in self.rows if getattr(r, "curso_id", None) == course_id), None)

    def _selected_attendee(self):
        selected = self.attendees_table.selectionModel().selectedRows()
        if not selected:
            return None
        item = self.attendees_table.item(selected[0].row(), 1)
        if item is None:
            return None
        contacto_id = str(item.data(Qt.ItemDataRole.UserRole) or "")
        return next((r for r in self.attendee_rows if r.contacto_id == contacto_id), None)

    def _selected_technician(self):
        selected = self.technicians_table.selectionModel().selectedRows()
        if not selected:
            return None
        item = self.technicians_table.item(selected[0].row(), 0)
        if item is None:
            return None
        tecnico_id = str(item.data(Qt.ItemDataRole.UserRole) or "")
        return next((r for r in self.technician_rows if r.tecnico_id == tecnico_id), None)

    def _schedule_reload(self) -> None:
        self._search_timer.start(250)

    def _schedule_autosave(self, *_args) -> None:
        if self._is_loading_details or not self._selected_course():
            return
        self._autosave_timer.start(350)

    def reload(self) -> None:
        year = self.year_filter.currentData()
        month_start = self.month_start_filter.currentData()
        month_end = self.month_end_filter.currentData()
        self.rows = self.service.list_courses(
            term=self.search_input.text().strip(),
            year=int(year) if year else None,
            month_start=int(month_start) if month_start else None,
            month_end=int(month_end) if month_end else None,
        )
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(self.rows))
        for i, row in enumerate(self.rows):
            date_cell = QTableWidgetItem(row.curso_fecha.strftime("%d/%m/%Y"))
            date_cell.setData(Qt.ItemDataRole.UserRole, row.curso_id)
            self.table.setItem(i, 0, date_cell)
            self.table.setItem(i, 1, QTableWidgetItem(str(row.curso_nombre or "")))
        self.table.setSortingEnabled(True)
        self._show_selected_details()

    def _show_selected_details(self) -> None:
        row = self._selected_course()
        self._is_loading_details = True
        if not row:
            self.detail_curso_nombre.clear()
            self.detail_curso_fecha.setDate(QDate.currentDate())
            self._render_attendees("")
            self._render_technicians("")
            self._set_documents_fields("", "", "")
            self._is_loading_details = False
            return
        self.detail_curso_nombre.setText(str(row.curso_nombre or ""))
        self.detail_curso_fecha.setDate(QDate(row.curso_fecha.year, row.curso_fecha.month, row.curso_fecha.day))
        docs = self.service.get_documents(row.curso_id)
        self._set_documents_fields(docs.portada, docs.invitacion, docs.recetario)
        self._render_attendees(row.curso_id)
        self._render_technicians(row.curso_id)
        self._is_loading_details = False

    def _render_attendees(self, curso_id: str) -> None:
        self._loading_attendees = True
        self.attendees_table.blockSignals(True)
        if not curso_id:
            self.attendee_rows = []
            self.attendees_table.setRowCount(0)
            self.attendees_table.blockSignals(False)
            self._loading_attendees = False
            self._update_attendees_counter()
            return
        self.attendee_rows = self.service.list_attendees(curso_id)
        rows = self._sorted_attendee_rows()
        self.attendees_table.setRowCount(len(rows))
        for i, row in enumerate(rows):
            confirm_cell = QTableWidgetItem("")
            confirm_cell.setFlags(Qt.ItemFlag.NoItemFlags)
            name_cell = QTableWidgetItem(row.asistente)
            name_cell.setData(Qt.ItemDataRole.UserRole, row.contacto_id)
            self.attendees_table.setItem(i, 0, confirm_cell)
            self.attendees_table.setItem(i, 1, name_cell)
            self.attendees_table.setItem(i, 2, QTableWidgetItem(row.nif))
            self.attendees_table.setItem(i, 3, QTableWidgetItem(row.empresa))
            self.attendees_table.setItem(i, 4, QTableWidgetItem(""))

            box = QCheckBox()
            box.setChecked(bool(row.status_confirmacion))
            box.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            box.setStyleSheet(
                "QCheckBox { background: transparent; margin: 0px; padding: 0px; }"
                "QCheckBox::indicator { width: 14px; height: 14px; margin: 0px; }"
            )
            box.toggled.connect(lambda checked, contacto_id=row.contacto_id: self._on_attendee_checkbox_toggled(contacto_id, checked))
            holder = QWidget()
            holder.setStyleSheet("background: transparent;")
            holder_layout = QHBoxLayout(holder)
            holder_layout.setContentsMargins(0, 0, 0, 0)
            holder_layout.setSpacing(0)
            holder_layout.addStretch(1)
            holder_layout.addWidget(box)
            holder_layout.addStretch(1)
            self.attendees_table.setCellWidget(i, 0, holder)
            self.attendees_table.setRowHeight(i, 35)

            edit_btn = QPushButton("")
            edit_btn.setProperty("btnRole", "success")
            edit_btn.setToolTip("Editar observaciones")
            edit_btn.setIcon(QIcon(str(PENCIL_ICON_PATH)))
            edit_btn.setIconSize(QSize(11, 11))
            edit_btn.setFlat(True)
            edit_btn.setStyleSheet(
                "QPushButton { min-height: 20px; max-height: 20px; padding: 0px 2px; "
                "background-color: #5BBE6A; border: none; border-radius: 4px; }"
                "QPushButton:hover { background-color: #49A85A; border: none; }"
                "QPushButton:pressed { background-color: #49A85A; border: none; }"
            )
            edit_btn.setFixedSize(22, 20)
            edit_btn.clicked.connect(
                lambda _checked=False, contacto_id=row.contacto_id: self._edit_attendee_observaciones(contacto_id)
            )
            btn_holder = QWidget()
            btn_holder.setStyleSheet("background: transparent; border: none;")
            btn_layout = QHBoxLayout(btn_holder)
            btn_layout.setContentsMargins(0, 0, 0, 0)
            btn_layout.setSpacing(0)
            btn_layout.addStretch(1)
            btn_layout.addWidget(edit_btn)
            btn_layout.addStretch(1)
            self.attendees_table.setCellWidget(i, 4, btn_holder)

        self.attendees_table.horizontalHeader().setSortIndicator(self._attendee_sort_col, self._attendee_sort_order)
        self.attendees_table.blockSignals(False)
        self._loading_attendees = False
        self._update_attendees_counter()

    def _update_attendees_counter(self) -> None:
        total = len(self.attendee_rows)
        confirmed = sum(1 for row in self.attendee_rows if bool(getattr(row, "status_confirmacion", False)))
        if hasattr(self, "attendees_counter_label"):
            self.attendees_counter_label.setText(f"{confirmed}/{total}")

    def _render_technicians(self, curso_id: str) -> None:
        if not curso_id:
            self.technician_rows = []
            self.technicians_table.setRowCount(0)
            return
        self.technician_rows = self.service.list_course_technicians(curso_id)
        self.technicians_table.setSortingEnabled(False)
        self.technicians_table.setRowCount(len(self.technician_rows))
        for i, row in enumerate(self.technician_rows):
            name_cell = QTableWidgetItem(row.nombre_completo)
            name_cell.setData(Qt.ItemDataRole.UserRole, row.tecnico_id)
            self.technicians_table.setItem(i, 0, name_cell)
            self.technicians_table.setItem(i, 1, QTableWidgetItem(row.movil))
            self.technicians_table.setItem(i, 2, QTableWidgetItem(row.interno))
            self.technicians_table.setItem(i, 3, QTableWidgetItem(row.email))
        self.technicians_table.setSortingEnabled(True)

    def _select_course_by_id(self, curso_id: str) -> None:
        for row in range(self.table.rowCount()):
            cell = self.table.item(row, 0)
            if cell and str(cell.data(Qt.ItemDataRole.UserRole) or "") == str(curso_id):
                self.table.selectRow(row)
                break

    def _on_attendee_item_changed(self, item: QTableWidgetItem) -> None:
        if self._loading_attendees or item.column() != 0:
            return
        # Check handling is done via checkbox widget, this remains for safety.

    def _on_attendee_checkbox_toggled(self, contacto_id: str, status: bool) -> None:
        if self._loading_attendees:
            return
        course = self._selected_course()
        if not course:
            return
        self.service.set_attendee_confirmation(course.curso_id, contacto_id, status)
        for row in self.attendee_rows:
            if row.contacto_id == contacto_id:
                row.status_confirmacion = status
                break
        self._update_attendees_counter()

    def _open_attendee_contact(self, row: int, _column: int) -> None:
        cell = self.attendees_table.item(row, 1)
        if not cell:
            return
        contacto_id = str(cell.data(Qt.ItemDataRole.UserRole) or "")
        if contacto_id:
            self._focus_contact_in_contacts_page(contacto_id)

    def _focus_contact_in_contacts_page(self, contacto_id: str) -> None:
        main_window = self.window()
        pages = getattr(main_window, "pages", None)
        if pages is None:
            return
        contacts_page = None
        for idx in range(pages.count()):
            widget = pages.widget(idx)
            if widget.__class__.__name__ == "ContactsPage":
                contacts_page = widget
                if hasattr(main_window, "_set_current_page"):
                    cast(Any, main_window)._set_current_page(idx)
                else:
                    pages.setCurrentIndex(idx)
                break
        if contacts_page is None:
            return
        contacts_page.reload()
        if hasattr(contacts_page, "_select_row_by_id"):
            contacts_page._select_row_by_id(contacto_id)

    def _edit_attendee_observaciones(self, contacto_id: str) -> None:
        course = self._selected_course()
        if not course:
            return
        attendee = next((r for r in self.attendee_rows if r.contacto_id == contacto_id), None)
        initial = attendee.observaciones if attendee else ""
        dialog = ObservacionesDialog("Observaciones", "Observaciones:", initial, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        notes = dialog.value()
        self.service.update_attendee_observaciones(course.curso_id, contacto_id, notes)
        if attendee:
            attendee.observaciones = notes.strip()

    def _sort_attendees_by_column(self, col: int) -> None:
        if col in {0, 4}:
            self.attendees_table.horizontalHeader().setSortIndicator(self._attendee_sort_col, self._attendee_sort_order)
            return
        if self._attendee_sort_col == col:
            self._attendee_sort_order = Qt.SortOrder.DescendingOrder if self._attendee_sort_order == Qt.SortOrder.AscendingOrder else Qt.SortOrder.AscendingOrder
        else:
            self._attendee_sort_col = col
            self._attendee_sort_order = Qt.SortOrder.AscendingOrder
        course = self._selected_course()
        if not course:
            return
        self._render_attendees(course.curso_id)

    def _sorted_attendee_rows(self):
        rows = list(self.attendee_rows)
        reverse = self._attendee_sort_order == Qt.SortOrder.DescendingOrder
        key_map = {
            1: lambda x: (x.asistente or "").lower(),
            2: lambda x: (x.nif or "").lower(),
            3: lambda x: (x.empresa or "").lower(),
        }
        key_fn = key_map.get(self._attendee_sort_col, key_map[1])
        rows.sort(key=key_fn, reverse=reverse)
        return rows

    @staticmethod
    def build_document_preview_widget(path_text: str, parent=None) -> QWidget:
        container = QWidget(parent)
        container_layout = QVBoxLayout(container)
        container_layout.setContentsMargins(0, 0, 0, 0)

        stack = QStackedWidget(container)
        placeholder = QLabel("Sin documento", stack)
        placeholder.setAlignment(Qt.AlignmentFlag.AlignCenter)
        stack.addWidget(placeholder)

        image_label = QLabel(stack)
        image_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        stack.addWidget(image_label)

        pdf_view = QPdfView(stack)
        pdf_doc = QPdfDocument(pdf_view)
        pdf_view.setDocument(pdf_doc)
        pdf_view.setPageMode(QPdfView.PageMode.MultiPage)
        pdf_view.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        stack.addWidget(pdf_view)

        text_view = QPlainTextEdit(stack)
        text_view.setReadOnly(True)
        stack.addWidget(text_view)
        container_layout.addWidget(stack, 1)

        setattr(container, "_doc_stack", stack)
        setattr(container, "_doc_placeholder", placeholder)
        setattr(container, "_doc_image", image_label)
        setattr(container, "_doc_pdf_view", pdf_view)
        setattr(container, "_doc_pdf_doc", pdf_doc)
        setattr(container, "_doc_text", text_view)
        CoursesPage.populate_document_preview_widget(container, path_text)
        return container

    @staticmethod
    def populate_document_preview_widget(preview_widget: QWidget, path_text: str) -> None:
        stack: QStackedWidget = cast(QStackedWidget, getattr(preview_widget, "_doc_stack"))
        placeholder: QLabel = cast(QLabel, getattr(preview_widget, "_doc_placeholder"))
        image_label: QLabel = cast(QLabel, getattr(preview_widget, "_doc_image"))
        pdf_doc: QPdfDocument = cast(QPdfDocument, getattr(preview_widget, "_doc_pdf_doc"))
        text_view: QPlainTextEdit = cast(QPlainTextEdit, getattr(preview_widget, "_doc_text"))

        if not path_text:
            placeholder.setText("Sin documento")
            stack.setCurrentWidget(placeholder)
            return
        path = Path(path_text)
        if not path.exists():
            placeholder.setText("Ruta no encontrada")
            stack.setCurrentWidget(placeholder)
            return
        suffix = path.suffix.lower()
        if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"}:
            pix = QPixmap(str(path))
            if pix.isNull():
                placeholder.setText("No se pudo cargar la imagen")
                stack.setCurrentWidget(placeholder)
                return
            image_label.setPixmap(pix.scaled(1600, 1200, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
            stack.setCurrentWidget(image_label)
            return
        if suffix == ".pdf":
            pdf_doc.load(str(path))
            stack.setCurrentWidget(cast(QWidget, getattr(preview_widget, "_doc_pdf_view")))
            return
        if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".log"}:
            try:
                text_view.setPlainText(path.read_text(encoding="utf-8", errors="replace"))
            except Exception:
                text_view.setPlainText("No se pudo cargar el archivo de texto.")
            stack.setCurrentWidget(text_view)
            return
        placeholder.setText(f"Documento cargado: {path.name}")
        stack.setCurrentWidget(placeholder)

    def _set_documents_fields(self, portada: str, invitacion: str, recetario: str) -> None:
        self.doc_paths["portada"] = str(portada or "")
        self.doc_paths["invitacion"] = str(invitacion or "")
        self.doc_paths["recetario"] = str(recetario or "")
        self._refresh_document_preview("portada")
        self._refresh_document_preview("invitacion")
        self._refresh_document_preview("recetario")

    def _autosave_selected_course(self) -> None:
        row = self._selected_course()
        if not row:
            return
        payload = {"curso_nombre": self.detail_curso_nombre.text().strip(), "curso_fecha": self.detail_curso_fecha.date().toPython()}
        if not payload["curso_nombre"]:
            return
        self.service.update_course(row.curso_id, payload)
        self.reload()

    def _new_course(self) -> None:
        schema = [
            {"name": "curso_id", "label": "Curso_ID"},
            {"name": "curso_nombre", "label": "Curso_Nombre"},
            {"name": "curso_fecha", "label": "Curso_Fecha", "default": date.today().isoformat()},
            {"name": "invitacion", "label": "Invitacion"},
            {"name": "portada", "label": "Portada"},
            {"name": "recetario", "label": "Recetario"},
        ]
        dialog = EntityDialog("Nuevo: Curso", schema, parent=self)
        if not dialog.exec():
            return
        try:
            self.service.create_course(dialog.get_payload())
        except Exception as exc:
            QMessageBox.warning(self, "Cursos", f"No se pudo crear el curso: {exc}")
        self.reload()

    def _delete_course(self) -> None:
        row = self._selected_course()
        if not row:
            QMessageBox.warning(self, "Cursos", "Selecciona un curso.")
            return
        if QMessageBox.question(self, "Confirmar", f"Eliminar curso {row.curso_nombre}?") != QMessageBox.StandardButton.Yes:
            return
        self.service.delete_course(row.curso_id)
        self.reload()

    def _import_courses(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(self, "Seleccionar archivo", "", "Archivos de datos (*.xlsx *.xlsm *.csv)")
        if not file_path:
            return
        schema = [
            {"name": "curso_id", "label": "Curso_ID"},
            {"name": "curso_nombre", "label": "Curso_Nombre"},
            {"name": "curso_fecha", "label": "Curso_Fecha"},
            {"name": "invitacion", "label": "Invitacion"},
            {"name": "portada", "label": "Portada"},
            {"name": "recetario", "label": "Recetario"},
        ]
        imported, errors = self.service.import_courses(Path(file_path))
        self.reload()
        if errors:
            QMessageBox.warning(self, "Importacion", f"Importados: {imported}\nErrores: {len(errors)}\n" + "\n".join(errors[:8]))
        else:
            QMessageBox.information(self, "Importacion", f"Importados: {imported}")

    def _add_attendee(self) -> None:
        course = self._selected_course()
        if not course:
            QMessageBox.warning(self, "Asistentes", "Selecciona un curso.")
            return
        picker = AttendeePickerDialog(self.service, self)
        if not picker.exec():
            return
        selected = picker.selected()
        if not selected:
            QMessageBox.warning(self, "Asistentes", "Selecciona un contacto.")
            return
        dialog = ObservacionesDialog("Observaciones", "Observaciones (opcional):", "", self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        obs = dialog.value()
        self.service.add_attendee(course.curso_id, selected.contacto_id, selected.cliente_id, obs.strip(), False)
        self.reload()
        self._select_course_by_id(course.curso_id)
        self._render_attendees(course.curso_id)

    def _delete_attendee(self) -> None:
        course = self._selected_course()
        attendee = self._selected_attendee()
        if not course or not attendee:
            QMessageBox.warning(self, "Asistentes", "Selecciona un asistente.")
            return
        self.service.remove_attendee(course.curso_id, attendee.contacto_id)
        self._render_attendees(course.curso_id)

    def _add_technician(self) -> None:
        course = self._selected_course()
        if not course:
            QMessageBox.warning(self, "Tecnicos", "Selecciona un curso.")
            return
        picker = TechnicianPickerDialog(self.service, self)
        if not picker.exec():
            return
        selected = picker.selected()
        if not selected:
            QMessageBox.warning(self, "Tecnicos", "Selecciona un tecnico.")
            return
        self.service.add_course_technician(course.curso_id, selected.tecnico_id)
        self._render_technicians(course.curso_id)

    def _delete_technician(self) -> None:
        course = self._selected_course()
        technician = self._selected_technician()
        if not course or not technician:
            QMessageBox.warning(self, "Tecnicos", "Selecciona un tecnico.")
            return
        self.service.remove_course_technician(course.curso_id, technician.tecnico_id)
        self._render_technicians(course.curso_id)

    def _import_attendees(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(self, "Seleccionar archivo", "", "Archivos de datos (*.xlsx *.xlsm *.csv)")
        if not file_path:
            return
        schema = [
            {"name": "curso_id", "label": "Curso_ID"},
            {"name": "contacto_id", "label": "Contacto_ID"},
            {"name": "cliente_id", "label": "Cliente_ID"},
            {"name": "observaciones", "label": "Observaciones"},
            {"name": "status_confirmacion", "label": "Status_Confirmacion", "type": "bool", "default": False},
        ]
        imported, errors = self.service.import_attendees(Path(file_path))
        selected_course = self._selected_course()
        if selected_course:
            selected_id = selected_course.curso_id
            self.reload()
            self._select_course_by_id(selected_id)
            self._render_attendees(selected_id)
        if errors:
            QMessageBox.warning(self, "Importacion", f"Importados: {imported}\nErrores: {len(errors)}\n" + "\n".join(errors[:8]))
        else:
            QMessageBox.information(self, "Importacion", f"Importados: {imported}")

    def _export_attendees_excel(self) -> None:
        course = self._selected_course()
        if not course:
            QMessageBox.warning(self, "Asistentes", "Selecciona un curso.")
            return
        save_path, _ = QFileDialog.getSaveFileName(self, "Guardar asistentes en Excel", "asistentes.xlsx", "Excel (*.xlsx)")
        if not save_path:
            return
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = "Asistentes"
        ws.append(["Confirmado", "Asistente", "NIF", "Empresa"])
        for row in self.attendee_rows:
            ws.append(["SI" if row.status_confirmacion else "NO", row.asistente, row.nif, row.empresa])
        wb.save(save_path)
        QMessageBox.information(self, "Asistentes", f"Archivo generado:\n{save_path}")

    def _export_attendees_pdf(self) -> None:
        course = self._selected_course()
        if not course:
            QMessageBox.warning(self, "Asistentes", "Selecciona un curso.")
            return
        save_path, _ = QFileDialog.getSaveFileName(self, "Guardar asistentes en PDF", "asistentes.pdf", "PDF (*.pdf)")
        if not save_path:
            return
        html_rows = "".join(
            f"<tr><td>{'SI' if r.status_confirmacion else 'NO'}</td><td>{html.escape(r.asistente)}</td><td>{html.escape(r.nif)}</td><td>{html.escape(r.empresa)}</td></tr>"
            for r in self.attendee_rows
        )
        html_doc = f"<h2>Asistentes del curso</h2><table border='1' cellspacing='0' cellpadding='4' width='100%'><thead><tr><th>Confirmado</th><th>Asistente</th><th>NIF</th><th>Empresa</th></tr></thead><tbody>{html_rows}</tbody></table>"
        writer = QPdfWriter(save_path)
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        writer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter)
        doc = QTextDocument()
        doc.setHtml(html_doc)
        painter = QPainter(writer)
        doc.drawContents(painter)
        painter.end()
        QMessageBox.information(self, "Asistentes", f"Archivo generado:\n{save_path}")

    def _generate_signature_pdf(self, scope: str, template_key: str) -> Path:
        course = self._selected_course()
        if not course:
            raise ValueError("Selecciona un curso.")
        return self.course_document_generation_service.generate_signature_pdf(
            course,
            self._sorted_attendee_rows(),
            scope=scope,
            template_key=template_key,
            selected_attendee=self._selected_attendee(),
        )

    def _preview_signature_sheets(self, scope: str, template_key: str) -> None:
        try:
            output_path = self._generate_signature_pdf(scope, template_key)
            dialog = DocumentExpandedDialog("Previsualizacion hojas de firma", str(output_path), self)
            dialog.exec()
        except Exception as exc:
            QMessageBox.warning(self, "Firmas", f"No se pudo previsualizar: {exc}")

    def _print_signature_sheets(self, scope: str, template_key: str) -> None:
        try:
            output_path = self._generate_signature_pdf(scope, template_key)
            self._print_pdf_file(str(output_path))
        except Exception as exc:
            QMessageBox.warning(self, "Firmas", f"No se pudo imprimir: {exc}")

    def _open_consentimientos_manager(self) -> None:
        dialog = ConsentimientosDialog(self)
        dialog.preview_btn.clicked.connect(
            lambda: self._preview_signature_sheets(dialog.selected_scope(), dialog.selected_template())
        )
        dialog.print_btn.clicked.connect(
            lambda: self._print_signature_sheets(dialog.selected_scope(), dialog.selected_template())
        )
        dialog.exec()

    def _generate_certificates_pdf(self, scope: str) -> Path:
        course = self._selected_course()
        if not course:
            raise ValueError("Selecciona un curso.")
        return self.course_document_generation_service.generate_certificates_pdf(
            course,
            self._sorted_attendee_rows(),
            scope=scope,
            selected_attendee=self._selected_attendee(),
        )

    def _preview_certificates(self, scope: str) -> None:
        try:
            output_path = self._generate_certificates_pdf(scope)
            dialog = DocumentExpandedDialog("Previsualizacion certificados", str(output_path), self)
            dialog.exec()
        except Exception as exc:
            QMessageBox.warning(self, "Certificados", f"No se pudo previsualizar: {exc}")

    def _print_certificates(self, scope: str) -> None:
        try:
            output_path = self._generate_certificates_pdf(scope)
            self._print_pdf_file(str(output_path), dialog_title="Certificados")
        except Exception as exc:
            QMessageBox.warning(self, "Certificados", f"No se pudo imprimir: {exc}")

    def _open_certificados_manager(self) -> None:
        dialog = CertificadosDialog(self)
        dialog.preview_btn.clicked.connect(lambda: self._preview_certificates(dialog.selected_scope()))
        dialog.print_btn.clicked.connect(lambda: self._print_certificates(dialog.selected_scope()))
        dialog.exec()

    def _attach_document(self, field_name: str) -> None:
        file_path, _ = QFileDialog.getOpenFileName(self, "Seleccionar archivo", "", "Todos los archivos (*.*)")
        if file_path:
            self.doc_paths[field_name] = file_path
            self._refresh_document_preview(field_name)
            self._save_documents(notify=False)

    def _refresh_document_preview(self, field_name: str) -> None:
        self.populate_document_preview_widget(self.doc_previews[field_name], self.doc_paths[field_name])

    def _delete_document(self, field_name: str) -> None:
        self.doc_paths[field_name] = ""
        self._refresh_document_preview(field_name)
        self._save_documents(notify=False)

    def _expand_document(self, field_name: str) -> None:
        title_map = {"portada": "Portada", "invitacion": "Invitación", "recetario": "Recetario"}
        dialog = DocumentExpandedDialog(
            f"Documento: {title_map.get(field_name, field_name)}",
            self.doc_paths.get(field_name, ""),
            self,
        )
        dialog.exec()

    def _print_document(self, field_name: str) -> None:
        path = self.doc_paths.get(field_name, "").strip()
        if not path:
            QMessageBox.warning(self, "Documentos", "No hay documento para imprimir.")
            return
        suffix = Path(path).suffix.lower()
        if suffix == ".pdf":
            self._print_pdf_file(path, dialog_title="Documentos")
            return
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        dialog = QPrintDialog(printer, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".log"}:
            doc = QTextDocument()
            try:
                doc.setPlainText(Path(path).read_text(encoding="utf-8", errors="replace"))
            except Exception:
                doc.setPlainText("No se pudo leer el archivo.")
            doc.print_(printer)
            return
        if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"}:
            pix = QPixmap(path)
            if pix.isNull():
                QMessageBox.warning(self, "Documentos", "No se pudo imprimir la imagen.")
                return
            painter = QPainter(printer)
            rect = painter.viewport()
            scaled = pix.scaled(rect.size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
            painter.drawPixmap((rect.width() - scaled.width()) // 2, (rect.height() - scaled.height()) // 2, scaled)
            painter.end()
            return
        QMessageBox.information(self, "Documentos", "Impresion no soportada para este formato.")

    def _print_pdf_file(self, path: str, dialog_title: str = "Firmas") -> None:
        # Use a short-lived document object and always close it to avoid file locking on Windows.
        pdf = QPdfDocument()
        try:
            load_error = pdf.load(path)
            if load_error != QPdfDocument.Error.None_:
                QMessageBox.warning(self, dialog_title, "No se pudo cargar el PDF para imprimir.")
                return
            printer = QPrinter(QPrinter.PrinterMode.HighResolution)
            dialog = QPrintDialog(printer, self)
            if dialog.exec() != QDialog.DialogCode.Accepted:
                return
            page_count = pdf.pageCount()
            painter = QPainter(printer)
            try:
                target_size_f = printer.pageRect(QPrinter.Unit.DevicePixel).size()
                target_size = QSize(max(1, int(target_size_f.width())), max(1, int(target_size_f.height())))
                for page_idx in range(page_count):
                    image = pdf.render(page_idx, target_size)
                    painter.drawImage(0, 0, image)
                    if page_idx < page_count - 1:
                        printer.newPage()
            finally:
                if painter.isActive():
                    painter.end()
        finally:
            pdf.close()

    def _save_documents(self, notify: bool = True) -> None:
        row = self._selected_course()
        if not row:
            if notify:
                QMessageBox.warning(self, "Documentos", "Selecciona un curso.")
            return
        payload = dict(self.doc_paths)
        self.service.save_documents(row.curso_id, payload)
        if notify:
            QMessageBox.information(self, "Documentos", "Documentos guardados.")

    def _clear_documents_fields(self) -> None:
        self._delete_document("portada")
        self._delete_document("invitacion")
        self._delete_document("recetario")



